from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

QUESTION_COLUMNS = {"问题", "question", "q", "faq问题", "用户问题"}
ANSWER_COLUMNS = {"答案", "answer", "a", "回复", "客服回答", "标准答案"}


def _normalize_header(value) -> str:
    return str(value or "").strip().lower().replace(" ", "")


def _load_excel_faq(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [{"chunks": [], "page_no": 1}]

    headers = [_normalize_header(value) for value in rows[0]]
    question_index = next((index for index, value in enumerate(headers) if value in QUESTION_COLUMNS), 0)
    answer_index = next((index for index, value in enumerate(headers) if value in ANSWER_COLUMNS), 1 if len(headers) > 1 else 0)

    chunks: list[dict] = []
    for row_number, row in enumerate(rows[1:], start=2):
        question = str(row[question_index] or "").strip() if question_index < len(row) else ""
        answer = str(row[answer_index] or "").strip() if answer_index < len(row) else ""
        if not question or not answer:
            continue
        chunks.append(
            {
                "content": f"问题：{question}\n答案：{answer}",
                "page_no": row_number,
                "question": question,
                "answer": answer,
            }
        )
    return [{"chunks": chunks, "page_no": 1}]


def load_document(path: str) -> list[dict]:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".xlsx":
        return _load_excel_faq(file_path)
    if suffix == ".pdf":
        reader = PdfReader(str(file_path))
        return [{"text": page.extract_text() or "", "page_no": index + 1} for index, page in enumerate(reader.pages)]
    if suffix in {".docx", ".doc"}:
        doc = Document(str(file_path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [{"text": text, "page_no": 1}]
    if suffix in {".txt", ".md"}:
        return [{"text": file_path.read_text(encoding="utf-8", errors="ignore"), "page_no": 1}]
    raise ValueError(f"不支持的文件类型：{suffix}")
